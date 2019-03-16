from PyPDF2 import PdfFileReader
import docx
from xlrd import open_workbook
from subprocess import Popen, PIPE
import csv
import xml.etree.ElementTree as ET
import os


def read_text_text(filepath, data_text):
    text_file = open(filepath, 'r')
    data_text = text_file.read()
    text_file.close()
    return data_text


def read_pdf_text(filepath, data_text):
    pdf_file_object = open(filepath, 'rb')
    pdf_file = PdfFileReader(pdf_file_object)
    for page in pdf_file.pages:
        data_text += page.extractText()
    return data_text


def read_docx_text(filepath, data_text):
    word_document = docx.Document(filepath)
    for paragraph in word_document.paragraphs:
        data_text += paragraph.text
    return data_text


def read_doc_text(filepath, data_text):
    popen_param = ['antiword', filepath]
    popen_output = Popen(popen_param, stdout=PIPE)
    stdout, stderr = popen_output.communicate()
    data_text += stdout.decode('ascii', 'ignore')
    return data_text


def read_doc_text_catdoc(filepath, data_text):
    fileopen = os.popen('catdoc -w "%s"' % filepath)
    data_text = fileopen.read()
    fileopen.close()
    return data_text
    #(fi, fo, fe) = os.popen3('catdoc -w "%s"' % filepath)
    #print(fi, fo, fe)
    # fi.close()
    #data_text = fo.read()
    # print(data_text)
    #erroroutput = fe.read()
    # fo.close()
    # fe.close()
    # if not erroroutput:
    # return data_text
    # else:
    # raise OSError("Executing the command caused an error: %s" %
    # erroroutput)


def read_excel_text(filepath, data_text):
    excel_file = open_workbook(filepath)
    excel_text = ''
    for sheet in excel_file.sheets():
        for row in range(sheet.nrows):
            for column in range(sheet.ncols):
                excel_text += str(sheet.cell(row, column).value)
                excel_text += '. '
            excel_text += '\n'
        excel_text += '\n'
    data_text = excel_text
    return data_text


def read_excel_rowtext(filepath, data_text):
    excel_file = open_workbook(filepath)
    excel_text = ''
    for sheet in excel_file.sheets():
        for row in range(sheet.nrows):
            for column in range(sheet.ncols):
                excel_text += str(sheet.cell(row, column).value)
                excel_text += '. '
            excel_text += '\n'
        excel_text += '\n'
    data_text = excel_text
    return data_text


def read_csv_text(filepath, data_text):
    with open(filepath) as csv_file:
        csv_dataread = csv.reader(csv_file, delimiter=',', quotechar=',')
        for row in csv_dataread:
            data_text += '. '.join(row)
            data_text += '\n'
    return data_text


def read_odt_text(filepath, data_text):
    popen_param = ['odt2txt', filepath]
    popen_output = Popen(popen_param, stdout=PIPE)
    stdout, stderr = popen_output.communicate()
    data_text += stdout.decode('ascii', 'ignore')
    return data_text


def read_xml_tree(filepath):
    tree = ET.parse(filepath)
    return tree